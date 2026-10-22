"""Prepare auditable FCL drafts from Zhengzhou XLSX and Chang'an PDF source tables.

Read-only extraction; this script never writes to the database. Every numeric source
cell becomes one container/equipment quote. Missing prices are not treated as zero.
"""
import argparse
import hashlib
import json
import pathlib
import re
import uuid
import openpyxl
import pdfplumber


def base(origin, destination, supplier, source, **values):
    q = dict(container='待确认', equipment='待确认', currency='', price=None,
             unit='柜', service='', via='', scope='站到站', validFrom='',
             validUntil='', notes='', fees=[], profitPreset={'mode':'markup','value':0})
    q.update(values)
    q['source'] = source
    return dict(id=str(uuid.uuid4()), origin=origin, destination=destination,
                supplier_id=supplier, status='draft', quote=q)


def source(path, sheet, cell, raw_route, raw_price):
    return dict(filename=path.name, sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
                sheet=sheet, cell=cell, rawRoute=raw_route, rawPrice=str(raw_price),
                quotationPeriod='2026年11月')


def zhengzhou(path, supplier):
    w = openpyxl.load_workbook(path, data_only=True)
    s = w['整柜价格方案']
    def value(row,col):
        cell=s.cell(row,col)
        for merged in s.merged_cells.ranges:
            if cell.coordinate in merged:
                return s.cell(merged.min_row,merged.min_col).value
        return cell.value
    result=[]
    options=[(7,'40尺','COC','柜'),(8,'40尺','SOC','柜'),(9,'20尺','COC','柜'),(10,'20尺','SOC','柜'),(11,'2×20尺','SOC','组')]
    for r in range(3,53):
        origin=str(value(r,5) or '').strip()
        raw_dest=str(value(r,6) or '').strip()
        if not origin or not raw_dest:
            raise ValueError(f'Missing route at row {r}')
        destination=re.sub(r'^(郑州|新郑)[-—]', '', raw_dest)
        service=' / '.join(label for label in ['快线','全时刻','南线','专列','半程'] if label in destination)
        currency={'美金':'USD','美元':'USD','人民币':'CNY'}.get(str(value(r,12) or '').strip(),'')
        notes=str(value(r,13) or '')
        for col,box,equipment,unit in options:
            price=s.cell(r,col).value
            if not isinstance(price,(int,float)):continue
            src=source(path,s.title,s.cell(r,col).coordinate,origin+' → '+raw_dest,price)
            src.update(rawCurrency=value(r,12),direction=value(r,3),lineGroup=value(r,2),periodBasis='文件名月份，具体有效期待确认')
            scope='国内段' if r==35 else '站到站'
            detail=notes+'\n报价月份来自文件名；具体有效期待确认。'
            if not currency:detail+='\n原表币种单元格为空，需确认币种。'
            result.append(base(origin,destination,supplier,src,container=box,equipment=equipment,
                               currency=currency,price=price,unit=unit,service=service,
                               scope=scope,notes=detail))
    return result


def parse_price(raw):
    compact=re.sub(r'\s+','',raw)
    currency='USD' if '$' in compact else 'CNY' if '¥' in compact else ''
    if '-X' in compact.upper():return currency,None,'公式报价，需提供X后确认；原公式保留。'
    match=re.search(r'[\d][\d,.]*',compact)
    if not match:return currency,None,'缺少可识别金额，待确认。'
    token=match.group()
    if re.fullmatch(r'\d{1,3}\.\d{3}',token):
        return currency,int(token.replace('.','')),'原文以点分隔千位（'+token+'），按千位整理，待核对。'
    return currency,float(token.replace(',','')),''


def route_parts(raw):
    clean=re.sub(r'\s+','',raw)
    service=' / '.join(label for label in ['公班','快线','时刻表','专班','定制','集采','散柜','增值'] if label in clean)
    clean=re.sub(r'[（(][^）)]*[）)]','',clean)
    clean=re.sub(r'(SOC|COC)(/.*)?$','',clean)
    parts=re.split('[-—]',clean)
    if len(parts)<2:raise ValueError('Route cannot be separated: '+raw)
    # The suffixes here are destination conditions, not extra railway legs.
    if '波季' in parts[1] and len(parts)>2:
        return parts[0],parts[1]+'（'+parts[-1]+'）','',service
    return parts[0],parts[-1],'/'.join(parts[1:-1]),service


def changan(path,supplier):
    result=[]
    with pdfplumber.open(path) as doc:
        for index in [0,1,2,3,4,6,7,8]:
            page=doc.pages[index]
            tables=page.extract_tables()
            if len(tables)!=1:raise ValueError(f'Unexpected tables on PDF page {index+1}')
            table=tables[0]
            text=page.extract_text() or ''
            notes=text.split('备注：',1)[-1] if '备注：' in text else ''
            group=''
            for n,row in enumerate(table[1:],2):
                if row[0] and row[0].startswith('备注'):break
                if row[0]:group=row[0].replace('\n','')
                if index>=6:
                    _,country,dest,raw=row
                    if not dest or not raw:continue
                    origin='西安（起运范围待确认）';via=group;service='境外分拨'
                    currency,price,normalization=parse_price(raw)
                    entries=[('待确认','待确认','柜',currency,price,normalization)]
                    scope='分拨报价（范围待确认）'
                    raw_route=group+' → '+dest
                else:
                    _,raw_route,raw=row
                    if not raw_route or not raw:continue
                    origin,dest,via,service=route_parts(raw_route)
                    scope='国内段' if '定制' in service or index==2 and '满洲里' in raw_route else '铁海联运' if index==1 else '范围待确认' if '境外段' in raw_route else '站到站'
                    if index==3 and ((origin in ['西安','阳平'] and dest in ['阿拉山口','霍尔果斯','果斯','山口']) or (dest in ['西安','阳平'] and origin in ['阿拉山口','霍尔果斯','果斯','山口'])):scope='国内段'
                    if index in [3,4] and (index==3 or '巴库' in raw_route or '/' in raw):
                        if '（40尺）' in raw:
                            parts=[('',None),('40尺',raw)]
                        else:parts=list(zip(['2×20尺','40尺','45尺'],raw.split('/')))
                        entries=[]
                        for box,part in parts:
                            if not part:continue
                            currency,price,normalization=parse_price(part)
                            if price is not None:entries.append((box,'待确认','组' if box=='2×20尺' else '柜',currency,price,normalization))
                    else:
                        currency,price,normalization=parse_price(raw)
                        equipment='COC' if 'COC' in raw_route else 'SOC' if 'SOC' in raw_route or index in [0,1] else '待确认'
                        box='20尺' if '/20尺' in raw_route else '40尺冷柜' if '/冷柜' in raw_route else '40尺' if index<=2 else '待确认'
                        entries=[(box,equipment,'柜',currency,price,normalization)]
                for box,equipment,unit,currency,price,normalization in entries:
                    src=source(path,'班列境外分拨价格' if index>=6 else '班列线路站点报价',f'第{index+1}页，表格第{n}行',raw_route,raw)
                    src.update(page=index+1,normalizationNotes=normalization,lineGroup=group)
                    detail=notes+'\n'+normalization+'\n客户版售价，最终采购成交价、适用箱型及附加费用需确认。'
                    if index in [0,1,6,7,8]:detail+='\n【第15页中欧报价说明】\n'+(doc.pages[14].extract_text() or '')
                    if index in [3,4]:detail+='\n【第16页相关线路说明】\n'+(doc.pages[15].extract_text() or '')
                    if index in [3,4]:detail+='\nSOC/COC未明确；有效期参考各线路附件，需确认。'
                    if index>=6:detail+='\n本页未明确完整起运范围、箱型与SOC/COC；仅按分拨报价归档，不能作为新增境外段费用直接叠加。'
                    if index==2:detail+='\n本页说明自2026年10月1日起适用，与11月文件版本一并保留，具体有效期需确认。'
                    result.append(base(origin,dest,supplier,src,container=box,equipment=equipment,currency=currency,
                                       price=price,unit=unit,service=service,via=via,scope=scope,
                                       validFrom='2026-11-01' if index in [0,1,6,7,8] else '',
                                       validUntil='2026-11-30' if index in [0,1,6,7,8] else '',notes=detail))
    return result


if __name__=='__main__':
    p=argparse.ArgumentParser()
    p.add_argument('--zhengzhou',type=pathlib.Path,required=True)
    p.add_argument('--changan',type=pathlib.Path,required=True)
    p.add_argument('--zz-supplier',required=True)
    p.add_argument('--xa-supplier',required=True)
    args=p.parse_args()
    records=zhengzhou(args.zhengzhou,args.zz_supplier)+changan(args.changan,args.xa_supplier)
    print(json.dumps(records,ensure_ascii=False))
